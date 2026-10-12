#!/usr/bin/env python3
"""Build a local, read-only Agristat snapshot from supplied OpenSTAT CSVs and RSBSA XLSX."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sqlite3
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


OPENSTAT = {
    "rice_corn_volume.csv": ("production_volume", "metric_tons"),
    "rice_corn_area_harvested.csv": ("area_harvested", "hectares"),
    "rice_corn_stock.csv": ("stock_inventory", "metric_tons"),
    "livestock_volume.csv": ("production_volume", "metric_tons"),
    "livestock_volume2.csv": ("production_volume", "metric_tons"),
    "poultry_volume.csv": ("production_volume", "metric_tons"),
    "chicken_volume.csv": ("production_volume", "metric_tons"),
    "major_crops_volume.csv": ("production_volume", "metric_tons"),
    "mango_banana_volume.csv": ("production_volume", "metric_tons"),
    "mango_banana_area_planted.csv": ("area_planted_or_harvested", "hectares"),
    "mango_banana_no_trees.csv": ("bearing_trees_hills_vines", "count"),
    "vegetables_root_crops_volume.csv": ("production_volume", "metric_tons"),
    "vegetables_root_crops_area.csv": ("area_planted_or_harvested", "hectares"),
}

PROFILE_GROUPS = (
    "registrants", "farmers", "farm_workers", "fisherfolk", "agri_youth",
    "rice_farmers", "corn_farmers", "hvcdp_farmers", "poultry_farmers", "livestock_farmers",
)
PROFILE_FIELDS = ("male", "female", "total", "average_age", "senior", "ip", "pwd", "four_ps", "arb")
AREA_FIELDS = (
    "physical_rice_area", "physical_corn_area", "declared_rice_area", "declared_corn_area",
    "verified_rice_area", "rice_verification_percent", "verified_corn_area", "corn_verification_percent",
)
MISSING = {"", "..", "...", "-", "na", "n/a"}


def number(value):
    if value is None or str(value).strip().lower() in MISSING:
        return None
    try:
        return float(str(value).replace(",", "").replace("%", "").strip())
    except ValueError as error:
        raise ValueError(f"Unexpected numeric value: {value!r}") from error


def integer(value):
    parsed = number(value)
    if parsed is None:
        return None
    if not parsed.is_integer():
        raise ValueError(f"Expected whole-number count: {value!r}")
    return int(parsed)


def period(header):
    match = re.fullmatch(r"(20\d{2})\s*(Annual|Semester\s*[12]|Quarter\s*[1-4]|Q\s*[1-4])", header.strip(), re.I)
    if not match:
        raise ValueError(f"Unrecognized period header: {header!r}")
    year, label = match.groups()
    label = re.sub(r"\s+", "", label).lower()
    if label == "annual":
        return int(year), "annual", None
    if label.startswith("semester"):
        return int(year), "semester", int(label[-1])
    return int(year), "quarter", int(label[-1])


def geography(raw):
    dots = len(raw) - len(raw.lstrip("."))
    if dots == 2:
        name = raw[2:].strip()
        return "region", "Western Visayas" if "western visayas" in name.lower() else name.title()
    if dots == 4:
        return "province", raw[4:].strip().title()
    raise ValueError(f"Unexpected OpenSTAT geography: {raw!r}")


SCHEMA = """
PRAGMA foreign_keys = ON;
CREATE TABLE source_file (
  id INTEGER PRIMARY KEY, filename TEXT NOT NULL UNIQUE, source_system TEXT NOT NULL,
  title TEXT NOT NULL, sha256 TEXT NOT NULL, imported_at_utc TEXT NOT NULL
);
CREATE TABLE openstat_observation (
  id INTEGER PRIMARY KEY, source_id INTEGER NOT NULL REFERENCES source_file(id),
  source_row INTEGER NOT NULL,
  measure TEXT NOT NULL, unit TEXT NOT NULL, commodity TEXT NOT NULL,
  geography_level TEXT NOT NULL CHECK(geography_level IN ('region','province','national')),
  geography TEXT NOT NULL, year INTEGER NOT NULL,
  period_type TEXT NOT NULL CHECK(period_type IN ('month','quarter','semester','annual')),
  period_number INTEGER, value REAL, raw_value TEXT NOT NULL,
  UNIQUE(source_id, commodity, geography_level, geography, year, period_type, period_number)
);
CREATE INDEX idx_openstat_lookup ON openstat_observation(measure, commodity, geography_level, geography, year, period_type);
CREATE TABLE rsbsa_geography (
  id INTEGER PRIMARY KEY, source_id INTEGER NOT NULL REFERENCES source_file(id),
  source_row INTEGER NOT NULL, level TEXT NOT NULL CHECK(level IN ('region','province','district','municipality')),
  province TEXT, district TEXT, municipality TEXT, source_label TEXT NOT NULL,
  UNIQUE(source_id, source_row)
);
CREATE TABLE rsbsa_profile (
  geography_id INTEGER NOT NULL REFERENCES rsbsa_geography(id), category TEXT NOT NULL,
  male REAL, female REAL, total REAL, average_age REAL,
  senior REAL, ip REAL, pwd REAL, four_ps REAL, arb REAL,
  PRIMARY KEY(geography_id, category)
);
CREATE TABLE rsbsa_area (
  geography_id INTEGER NOT NULL REFERENCES rsbsa_geography(id),
  measure TEXT NOT NULL, value REAL, unit TEXT NOT NULL,
  PRIMARY KEY(geography_id, measure)
);
CREATE TABLE rsbsa_region_rank (
  source_id INTEGER NOT NULL REFERENCES source_file(id), source_row INTEGER NOT NULL,
  region TEXT NOT NULL, registrants INTEGER, registrants_rank INTEGER,
  verified_area_ha REAL, verified_area_rank INTEGER,
  PRIMARY KEY(source_id, source_row)
);
"""


def add_source(db, path, system, title):
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    cursor = db.execute(
        "INSERT INTO source_file(filename,source_system,title,sha256,imported_at_utc) VALUES (?,?,?,?,?)",
        (path.name, system, title, digest, datetime.now(timezone.utc).isoformat()),
    )
    return cursor.lastrowid


def import_openstat(db, folder, report):
    for filename, (measure, unit) in OPENSTAT.items():
        path = folder / filename
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle, delimiter=";"))
        if len(rows) < 4:
            raise ValueError(f"{filename}: too few rows")
        title = rows[0][0].strip()
        headers = rows[2]
        is_stock = filename == "rice_corn_stock.csv"
        source_id = add_source(db, path, "PSA OpenSTAT CSV export", title)
        count = missing = 0
        for source_row, row in enumerate(rows[3:], 4):
            if not any(cell.strip() for cell in row):
                continue
            if len(row) != len(headers):
                raise ValueError(f"{filename}:{source_row}: {len(row)} cells; expected {len(headers)}")
            commodity = row[0].strip()
            if is_stock:
                geo_level, geo_name, year = "national", "Philippines", int(row[1])
            else:
                geo_level, geo_name = geography(row[1])
            for col, raw in enumerate(row[2:], 2):
                if is_stock:
                    month = datetime.strptime(headers[col].strip(), "%B").month
                    p_year, p_type, p_number = year, "month", month
                else:
                    p_year, p_type, p_number = period(headers[col])
                value = number(raw)
                missing += value is None
                db.execute(
                    """INSERT INTO openstat_observation
                    (source_id,source_row,measure,unit,commodity,geography_level,geography,year,period_type,period_number,value,raw_value)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (source_id, source_row, measure, unit, commodity, geo_level, geo_name, p_year, p_type, p_number, value, raw.strip()),
                )
                count += 1
        report[filename] = {"observations": count, "missing": missing, "measure": measure, "unit": unit}


def classify_profile(row):
    province = str(row[1] or "").strip()
    place = str(row[2] or "").strip()
    district = str(row[3] or "").strip()
    if province.upper() == "TOTAL":
        return "region", None, None, None, "Western Visayas total"
    if not place:
        return None
    if re.fullmatch(r"\d+", place):
        if district.lower().startswith(("district", "distric", "lone district")):
            return "district", province, district, None, district
        return "province", province, None, None, province
    return "municipality", province, None, place.replace(" Capital", "").strip(), place


def import_rsbsa(db, path, report):
    book = load_workbook(path, read_only=True, data_only=True)
    if "Quick Database" not in book.sheetnames:
        raise ValueError("RSBSA workbook lacks 'Quick Database'")
    source_id = add_source(db, path, "RSBSA aggregate workbook", "RSBSA Data (09-18-26) / Quick Database")
    sheet = book["Quick Database"]
    header = list(sheet.iter_rows(min_row=3, max_row=3, values_only=True))[0]
    if str(header[4]).strip() != "Male" or str(header[6]).strip() != "Total":
        raise ValueError("RSBSA profile columns changed")
    levels = Counter()
    count_mismatches = []
    fractional_counts = []
    current_district = None
    current_province = None
    for source_row, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), 4):
        if not any(value is not None for value in row[1:7]):
            continue
        classification = classify_profile(row)
        if classification is None:
            continue
        level, province, district, municipality, label = classification
        if level == "province":
            current_province, current_district = province, None
        elif level == "district":
            current_province, current_district = province, district
        elif level == "municipality":
            district = current_district if province == current_province else None
        geo_id = db.execute(
            "INSERT INTO rsbsa_geography(source_id,source_row,level,province,district,municipality,source_label) VALUES (?,?,?,?,?,?,?)",
            (source_id, source_row, level, province, district, municipality, label),
        ).lastrowid
        levels[level] += 1
        for group_index, group in enumerate(PROFILE_GROUPS):
            start = 4 + group_index * 9
            raw = row[start:start + 9]
            values = [number(value) for value in raw]
            if all(value is None for value in values):
                continue
            counts = values
            db.execute(
                "INSERT INTO rsbsa_profile VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                (geo_id, group, *counts),
            )
            if all(value is not None for value in counts[:3]) and abs(counts[0] + counts[1] - counts[2]) > 0.000001:
                count_mismatches.append({"row": source_row, "category": group, "male_plus_female": counts[0] + counts[1], "total": counts[2]})
            for field, value in zip(PROFILE_FIELDS, counts):
                if field != "average_age" and value is not None and not value.is_integer():
                    fractional_counts.append({"row": source_row, "category": group, "field": field, "value": value})
        for offset, measure in enumerate(AREA_FIELDS):
            value = number(row[94 + offset])
            if value is not None:
                db.execute("INSERT INTO rsbsa_area VALUES (?,?,?,?)", (geo_id, measure, value, "percent" if measure.endswith("percent") else "hectares"))
    rank_count = 0
    if "Western Visayas Rank" in book.sheetnames:
        for source_row, row in enumerate(book["Western Visayas Rank"].iter_rows(min_row=4, values_only=True), 4):
            region = str(row[1] or "").strip()
            if not region or region.upper() == "PHILIPPINES":
                continue
            if number(row[2]) is None and number(row[4]) is None:
                continue
            db.execute(
                "INSERT INTO rsbsa_region_rank VALUES (?,?,?,?,?,?,?)",
                (source_id, source_row, region, integer(row[2]), integer(row[3]), number(row[4]), integer(row[5])),
            )
            rank_count += 1
    book.close()
    report["rsbsa"] = {"geographies": dict(levels), "region_rank_rows": rank_count,
                       "sex_total_mismatch_count": len(count_mismatches), "sex_total_mismatch_examples": count_mismatches[:10],
                       "fractional_count_count": len(fractional_counts), "fractional_count_examples": fractional_counts[:10]}


def audits(db, report):
    overview = {}
    overview["openstat_observations"] = db.execute("SELECT COUNT(*) FROM openstat_observation").fetchone()[0]
    overview["openstat_nonmissing"] = db.execute("SELECT COUNT(*) FROM openstat_observation WHERE value IS NOT NULL").fetchone()[0]
    overview["rsbsa_profile_rows"] = db.execute("SELECT COUNT(*) FROM rsbsa_profile").fetchone()[0]
    overview["rsbsa_area_rows"] = db.execute("SELECT COUNT(*) FROM rsbsa_area").fetchone()[0]
    detail = db.execute("""SELECT p.total FROM rsbsa_profile p JOIN rsbsa_geography g ON p.geography_id=g.id
                          WHERE g.level='region' AND p.category='registrants'""").fetchone()
    rank = db.execute("SELECT registrants FROM rsbsa_region_rank WHERE region='REGION VI'").fetchone()
    overview["rsbsa_regional_total"] = detail[0] if detail else None
    overview["rsbsa_rank_tab_region_vi"] = rank[0] if rank else None
    overview["rsbsa_rank_difference"] = detail[0] - rank[0] if detail and rank else None
    province_sum = db.execute("""SELECT SUM(p.total) FROM rsbsa_profile p JOIN rsbsa_geography g ON p.geography_id=g.id
                                WHERE g.level='province' AND p.category='registrants'""").fetchone()[0]
    muni_sum = db.execute("""SELECT SUM(p.total) FROM rsbsa_profile p JOIN rsbsa_geography g ON p.geography_id=g.id
                            WHERE g.level='municipality' AND p.category='registrants'""").fetchone()[0]
    overview["rsbsa_province_sum"] = province_sum
    overview["rsbsa_municipality_sum"] = muni_sum
    report["overview"] = overview
    report["overlapping_openstat_exports"] = [
        {"left": left, "right": right, "matching_nonmissing_cells": matches,
         "different_cells_over_0_1": different}
        for left, right, matches, different in db.execute("""
            SELECT a.filename,b.filename,COUNT(*),SUM(ABS(x.value-y.value)>0.1)
            FROM openstat_observation x JOIN openstat_observation y
              ON x.id<y.id AND x.measure=y.measure AND x.commodity=y.commodity
              AND x.geography_level=y.geography_level AND x.geography=y.geography
              AND x.year=y.year AND x.period_type=y.period_type
              AND x.period_number IS y.period_number
            JOIN source_file a ON a.id=x.source_id JOIN source_file b ON b.id=y.source_id
            WHERE x.value IS NOT NULL AND y.value IS NOT NULL
            GROUP BY a.filename,b.filename ORDER BY COUNT(*) DESC
        """)
    ]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, required=True, help="Folder containing the 13 supplied OpenSTAT CSVs")
    parser.add_argument("--rsbsa", type=Path, required=True, help="Supplied RSBSA workbook")
    parser.add_argument("--output", type=Path, default=Path("data/agristat.sqlite"))
    args = parser.parse_args()
    missing = sorted(set(OPENSTAT) - {path.name for path in args.source_dir.glob("*.csv")})
    if missing:
        parser.error(f"Missing CSVs: {', '.join(missing)}")
    if not args.rsbsa.is_file():
        parser.error(f"Missing RSBSA workbook: {args.rsbsa}")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    temporary = args.output.with_suffix(args.output.suffix + ".tmp")
    temporary.unlink(missing_ok=True)
    report = {"sources": {}}
    try:
        with sqlite3.connect(temporary) as db:
            db.executescript(SCHEMA)
            import_openstat(db, args.source_dir, report["sources"])
            import_rsbsa(db, args.rsbsa, report["sources"])
            audits(db, report)
            integrity = db.execute("PRAGMA integrity_check").fetchone()[0]
            if integrity != "ok":
                raise ValueError(f"SQLite integrity check: {integrity}")
        temporary.replace(args.output)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    report_path = args.output.with_suffix(".report.json")
    report_path.write_text(json.dumps(report, indent=2) + "\n")
    print(f"Built {args.output} and {report_path}")
    print(json.dumps(report["overview"], indent=2))


if __name__ == "__main__":
    main()
